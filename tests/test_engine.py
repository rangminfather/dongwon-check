"""대사 엔진 테스트.

합성 데이터로 양식별 인식·판정 규칙을 확인하고,
DONGWON_SAMPLES 환경변수에 실제 자료 폴더가 있으면 기존 수작업 결과(동원원장_비교.xlsx)와 대조한다.
실제 자료는 거래처 정보가 들어 있으므로 저장소에 올리지 않는다.
"""

import datetime as dt
import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

from dongwon_check import engine as E  # noqa: E402
from dongwon_check import exporter  # noqa: E402

D = lambda d: dt.datetime(2026, 9, d)  # noqa: E731

LEDGER_HEAD = ["납품요청일", "판매문서", "판매처", "판매처명", "바코드", "자재 내역", "총 수량", "단가", "공급가", "지점명"]


def ledger_rows():
    return [
        LEDGER_HEAD,
        [D(1), "1", "100", "가마트", "18801047305368", "개성왕만두 2.1kg", 4, 13804, 55216, "대구1지점"],  # 박스코드
        [D(1), "1", "100", "가마트", "8801047511267", "가쓰오크랩스", 3, 2236, 6708, "대구1지점"],
        [D(1), "1", "100", "가마트", "8801047511267", "가쓰오크랩스", 2, 2236, 4472, "대구1지점"],  # 같은 날 합산
        [D(2), "2", "100", "가마트", "8801047529002", "얇은사각", 5, 1658, 8290, "대구1지점"],
        [D(4), "3", "100", "가마트", "8801047513162", "싱싱맛살", -2, 2350, -4700, "대구1지점"],  # 반품, 거래처는 9/2
        [D(5), "4", "100", "가마트", "8801047529545", "갈치사각", 5, 1691, 8455, "대구1지점"],  # 단가차이
        [D(6), "5", "100", "가마트", "8801047557616", "풍미사각", 3, 2895, 8685, "대구1지점"],  # 동원만
        [D(6), "6", "100", "가마트", "8801047000000", "판촉물", 0, 0, 0, "대구1지점"],  # 0행 제외
        [D(1), "7", "200", "나마트", "8801047999990", "다른 판매처", 9, 100, 900, "대구1지점"],
    ]


def vendor_footer_rows():
    """양식2형: 품목 뒤에 '[MM월 DD일 HH:MM 합계]' 꼬리행."""
    head = ["코드", "공급사", "상품명", "규격", "바코드", "매입량", "공급가액", "부가세", "보증금", "매입액", "매입단가"]
    t = lambda d: ["", f"[09월 {d:02d}일 13:53 합계]", "", "", "", "", 0, 0, 0, 0, 0]  # noqa: E731
    return [
        head,
        ["1", "(주)동원F&B", "개성왕만두", "2100g", "8801047305361", 4, 55216, 5521.6, 0, 60737.6, 15184.4],
        ["1", "(주)동원F&B", "가쓰오크랩스", "144g", "8801047511267", 5, 11180, 1118, 0, 12298, 2459.6],
        t(1),
        ["1", "(주)동원F&B", "얇은사각", "200g", "8801047529002", 4, 6632, 663, 0, 7295, 1823.8],  # 수량차이
        ["1", "(주)동원F&B", "싱싱맛살", "", "8801047513162", -2, -4700, -470, 0, -5170, 2585],
        ["1", "농협", "다른회사 상품", "", "8809999999990", 1, 1000, 100, 0, 1100, 1100],  # 동원 외
        t(2),
        ["1", "(주)동원F&B", "갈치사각", "180g", "8801047529545", 5, 8000, 800, 0, 8800, 1760],
        ["1", "(주)동원F&B", "새상품", "", "8801047111110", 2, 3000, 300, 0, 3300, 1650],  # 거래처만
        t(5),
    ]


def vendor_grouped_rows():
    """양식1형: '매입일자 : YYYY-MM-DD' 머리행 + 헤더와 데이터 열이 어긋난 병합 헤더."""
    return [
        ["바코드", None, None, "상품명", "규격", "수량", "매입단가", "부가세액", "공급가액", "매입금액", "거래처"],
        ["매입일자 : 2026-09-01"],
        [None, "전표번호 : 1"],
        [None, None, "8801047305361", "개성왕만두", None, 4, 15184, 5521, 55216, 60737, "동원"],
        ["매입일자 : 2026-09-02"],
        [None, None, "8801047529002", "얇은사각", None, 5, 1823, 829, 8290, 9119, "동원"],
    ]


def vendor_period_rows():
    """양식3형: 날짜 없음, 매입금액(부가세 포함) − 부가세액."""
    return [
        [None, "상품코드", "상품명", "규격", "유형", "매입수량", "매입단가", "매입금액", "부가세액"],
        ["3", "8801047511267", "가쓰오크랩스", "144g", "과세", 5, 2459.6, 12298, 1118],
        ["1", "8801047529002", "얇은사각", "200g", "과세", 5, 1823.8, 9119, 829],
    ]


def test_barcode_key_box_and_ean():
    assert E.barcode_key("18801047305368") == E.barcode_key("8801047305361") == "880104730536"
    assert E.display_barcode("880104730536") == "8801047305361"


def test_ledger_parse_skips_zero_and_reads_branch():
    L = E.parse_ledger(ledger_rows(), "Data")
    assert len(L.records) == 8
    assert L.records[0].extra["branch"] == "대구1지점"


def test_footer_format_and_reconcile():
    L = E.parse_ledger(ledger_rows())
    V = E.parse_vendor(vendor_footer_rows(), period=(2026, 9))
    assert any("동원 외" in n for n in V.notes)
    res = E.compare(L, V, label="t")
    assert res.customer == "가마트"  # 겹침이 많은 판매처 자동 선택
    by = {(r.barcode, r.date): r for r in res.rows}
    assert by[("8801047305361", dt.date(2026, 9, 1))].status == "일치"
    gas = by[("8801047511267", dt.date(2026, 9, 1))]
    assert (gas.l_qty, gas.status) == (5, "일치")  # 같은 날 2행 합산
    assert by[("8801047529002", dt.date(2026, 9, 2))].status == "수량차이"
    ret = by[("8801047513162", dt.date(2026, 9, 4))]
    assert ret.status == "일치" and ret.v_date == dt.date(2026, 9, 2) and ret.memo == "거래처 일자 -2일"
    assert by[("8801047529545", dt.date(2026, 9, 5))].status == "단가차이"
    assert by[("8801047557616", dt.date(2026, 9, 6))].status == "동원만"
    new = next(r for r in res.rows if r.key == "880104711111")
    assert (new.date, new.status) == (dt.date(2026, 9, 5), "거래처만")


def test_grouped_format_dates_and_misaligned_header():
    V = E.parse_vendor(vendor_grouped_rows())
    assert [(r.date, r.qty, r.amount) for r in V.records] == [
        (dt.date(2026, 9, 1), 4, 55216), (dt.date(2026, 9, 2), 5, 8290)]


def test_period_format_uses_gross_minus_vat():
    L = E.parse_ledger(ledger_rows())
    V = E.parse_vendor(vendor_period_rows())
    assert V.records[0].amount == pytest.approx(11180)
    res = E.compare(L, V, label="p", customer="가마트")
    assert res.mode == "period"
    gas = next(r for r in res.rows if r.barcode == "8801047511267")
    assert gas.status == "일치" and gas.date is None


def test_window_zero_disables_pairing():
    L = E.parse_ledger(ledger_rows())
    V = E.parse_vendor(vendor_footer_rows(), period=(2026, 9))
    res = E.compare(L, V, label="t", window_days=0)
    assert not any(r.v_date for r in res.rows)


def test_export_roundtrip(tmp_path):
    L = E.parse_ledger(ledger_rows(), "Data")
    V = E.parse_vendor(vendor_footer_rows(), "거래처", period=(2026, 9))
    res = E.compare(L, V, label="양식2")
    out = exporter.export(str(tmp_path / "out.xlsx"), L, "원장.xlsx", [res])
    wb = openpyxl.load_workbook(out, data_only=True)
    assert wb.sheetnames == ["동원", "거래처", "비교"]
    ws = wb["비교"]
    assert ws["L12"].value in E.STATUSES  # 계산값이 함께 저장되어 있어야 한다
    out2 = exporter.export(str(tmp_path / "multi.xlsx"), L, "원장.xlsx", [res, E.compare(L, V, label="양식2")])
    assert openpyxl.load_workbook(out2).sheetnames[:3] == ["요약", "동원", "비교_양식2"]


SAMPLES = os.environ.get("DONGWON_SAMPLES")


@pytest.mark.skipif(not SAMPLES, reason="실제 자료 폴더(DONGWON_SAMPLES)가 없습니다")
def test_matches_manual_result():
    sheets = E.load_workbook_rows(os.path.join(SAMPLES, "동원원장_비교.xlsx"))
    L = E.parse_ledger(sheets["동원"])
    V = E.parse_vendor(sheets["거래처"], period=(2026, 9))
    res = E.compare(L, V, label="x")
    assert res.summary.counts == {"일치": 235, "수량차이": 91, "단가차이": 239, "동원만": 140, "거래처만": 76}
    assert res.summary.l_amt == pytest.approx(30372613)
    assert res.summary.v_amt == pytest.approx(30243213.6)
    assert res.warning == ""  # 같은 매장 자료는 경고 없음


def test_affiliation_detection():
    sheets = {"매입상세내역": vendor_grouped_rows(), "2609 영해식자재 잔액": [[None]]}
    aff = E.detect_affiliation("양식4.xlsx", sheets)
    assert (aff.store, aff.store_source) == ("영해식자재", "'2609 영해식자재 잔액' 시트 이름")
    assert E.detect_affiliation("양식1.xlsx", {"x": vendor_grouped_rows()}).store == ""  # 일반적인 파일 이름은 무시
    assert E.same_store("세종경산", "(주)세종마트 경산점")
    assert not E.same_store("영해식자재", "주식회사 지마트유통센터")
    L = E.parse_ledger(ledger_rows())
    V = E.parse_vendor(vendor_footer_rows(), period=(2026, 9))
    res = E.compare(L, V, label="t", affiliation=aff)
    assert "영해식자재" in res.warning and res.affiliation.supplier == "(주)동원F&B"
