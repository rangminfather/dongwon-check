"""동원 원장 ↔ 거래처 매입자료 대사(對査) 엔진.

UI 와 분리된 순수 로직이다. 흐름은
    load_workbook_rows → parse_ledger / parse_vendor → reconcile → (exporter)
이며, 거래처 양식은 열 이름·날짜 표기가 제각각이라 "헤더 별칭 + 내용 추론"으로 읽는다.
"""

from __future__ import annotations

import datetime as dt
import os
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field

# ---------------------------------------------------------------------------
# 파일 읽기
# ---------------------------------------------------------------------------


def load_workbook_rows(path: str) -> dict[str, list[list]]:
    """통합문서의 모든 시트를 {시트명: 행 목록} 으로 읽는다(.xlsx/.xlsm/.xls)."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xls":
        import xlrd

        book = xlrd.open_workbook(path)
        sheets = {}
        for sh in book.sheets():
            rows = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                    elif cell.ctype == xlrd.XL_CELL_EMPTY:
                        row.append(None)
                    else:
                        row.append(cell.value)
                rows.append(row)
            sheets[sh.name] = rows
        return sheets

    import openpyxl

    # read_only 로 열어야 이미지가 큰 파일(수십 MB)도 빠르게 읽힌다.
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# 값 정규화
# ---------------------------------------------------------------------------

_BARCODE_RE = re.compile(r"^\d{8,14}$")


def norm_header(v) -> str:
    return re.sub(r"\s+", "", str(v)) if v is not None else ""


def as_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def barcode_text(v) -> str | None:
    """셀 값이 바코드처럼 보이면 숫자 문자열로, 아니면 None."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and not v.is_integer():
            return None
        s = str(int(v))
    else:
        s = str(v).strip().replace("-", "")
    return s if _BARCODE_RE.match(s) else None


def ean_check_digit(body12: str) -> str:
    total = sum(int(d) * (3 if i % 2 else 1) for i, d in enumerate(body12))
    return str((10 - total % 10) % 10)


def barcode_key(code: str) -> str:
    """매칭 키 = 체크숫자를 뺀 본체 12자리.

    박스코드(ITF-14)는 앞 물류지시자 1자리와 끝 체크숫자를, EAN-13 은 끝 체크숫자를 뗀다.
    """
    if len(code) == 14:
        return code[1:13]
    if len(code) == 13:
        return code[:12]
    if len(code) == 12:  # UPC-A → EAN-13 형태로 맞춘다
        return ("0" + code)[:12]
    return code


def display_barcode(key: str) -> str:
    return key + ean_check_digit(key) if len(key) == 12 and key.isdigit() else key


def to_number(v) -> float | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if s in ("", "-"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


_YMD_RE = re.compile(r"(\d{4})\s*[-./년]\s*(\d{1,2})\s*[-./월]\s*(\d{1,2})")


def to_date(v) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str):
        m = _YMD_RE.search(v)
        if m:
            try:
                return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
            except ValueError:
                return None
    return None


# ---------------------------------------------------------------------------
# 공통 레코드
# ---------------------------------------------------------------------------


@dataclass
class Record:
    date: dt.date | None
    key: str
    barcode: str
    name: str
    qty: float
    amount: float  # 공급가(부가세 제외)
    row: int  # 원본 엑셀 행 번호(1부터)
    customer: str = ""
    extra: dict = field(default_factory=dict)


@dataclass
class ParseResult:
    kind: str  # "ledger" | "vendor"
    sheet: str
    records: list[Record]
    columns: dict[str, int]  # 인식한 열(0부터)
    header_row: int  # 1부터
    notes: list[str] = field(default_factory=list)
    raw_header: list = field(default_factory=list)
    raw_rows: list[list] = field(default_factory=list)
    raw_all: list[list] = field(default_factory=list)  # 시트 전체(원본행 번호 = 인덱스+1)


ALIASES: dict[str, list[str]] = {
    "barcode": ["바코드", "상품바코드", "상품코드", "barcode", "EAN"],
    "name": ["상품명", "자재내역", "자재명", "품명", "제품명", "품목명"],
    "spec": ["규격"],
    "qty": ["총수량", "매입수량", "매입량", "입고수량", "수량"],
    "supply": ["공급가액", "공급가", "공급금액"],
    "gross": ["매입금액", "매입액", "합계금액", "총액", "금액"],
    "vat": ["부가세액", "부가세"],
    "date": ["납품요청일", "매입일자", "입고일자", "납품일자", "거래일자", "청구일", "일자"],
    "supplier": ["공급사", "거래처", "매입처"],
    "customer": ["판매처명"],
    "customer_code": ["판매처"],
    "branch": ["지점명"],
    "unit": ["매입단가", "단가"],
}


def map_header(header: list) -> dict[str, int]:
    norm = [norm_header(h) for h in header]
    found: dict[str, int] = {}
    for field_name, names in ALIASES.items():
        for alias in names:
            a = norm_header(alias)
            if a in norm:
                found[field_name] = norm.index(a)
                break
    return found


# ---------------------------------------------------------------------------
# 동원 원장
# ---------------------------------------------------------------------------


def looks_like_ledger(rows: list[list]) -> bool:
    for row in rows[:5]:
        hs = {norm_header(v) for v in row}
        if "바코드" in hs and ("총수량" in hs or "공급가" in hs) and ("판매처명" in hs or "자재내역" in hs):
            return True
    return False


def parse_ledger(rows: list[list], sheet: str = "") -> ParseResult:
    hdr_idx = next(
        i for i, row in enumerate(rows[:5]) if {"바코드"} <= {norm_header(v) for v in row}
    )
    header = rows[hdr_idx]
    cols = map_header(header)
    for need in ("barcode", "qty", "supply", "date"):
        if need not in cols:
            raise ValueError(f"원장에서 '{need}' 열을 찾지 못했습니다.")
    recs: list[Record] = []
    skipped = 0
    data_rows = rows[hdr_idx + 1 :]
    for i, row in enumerate(data_rows, start=hdr_idx + 2):
        g = lambda f: row[cols[f]] if f in cols and cols[f] < len(row) else None  # noqa: E731
        code = barcode_text(g("barcode"))
        if not code:
            continue
        qty = to_number(g("qty")) or 0.0
        amt = to_number(g("supply")) or 0.0
        if qty == 0 and amt == 0:
            skipped += 1
            continue
        recs.append(
            Record(
                date=to_date(g("date")),
                key=barcode_key(code),
                barcode=code,
                name=as_text(g("name")),
                qty=qty,
                amount=amt,
                row=i,
                customer=as_text(g("customer")) or as_text(g("customer_code")),
                extra={"branch": as_text(g("branch"))},
            )
        )
    notes = [f"원장 {len(recs):,}행 인식"]
    if skipped:
        notes.append(f"수량·금액 0 행 {skipped}건 제외")
    return ParseResult("ledger", sheet, recs, cols, hdr_idx + 1, notes, list(header), data_rows, rows)


# ---------------------------------------------------------------------------
# 거래처 매입자료 (양식 자동 인식)
# ---------------------------------------------------------------------------

_GROUP_DATE_RE = re.compile(r"일자\s*[:：]\s*(\d{4})\s*[-./]\s*(\d{1,2})\s*[-./]\s*(\d{1,2})")
_FOOTER_RE = re.compile(r"(\d{1,2})\s*월\s*(\d{1,2})\s*일[^\]]*합계")


def _find_header_row(rows: list[list]) -> tuple[int, dict[str, int]] | None:
    best = None
    for i, row in enumerate(rows[:40]):
        cols = map_header(row)
        if "qty" not in cols or not ({"supply", "gross"} & cols.keys()):
            continue
        score = len(cols)
        if best is None or score > best[0]:
            best = (score, i, cols)
    return (best[1], best[2]) if best else None


def _barcode_column(rows: list[list], start: int, hinted: int | None) -> int | None:
    counts: Counter = Counter()
    for row in rows[start : start + 400]:
        for j, v in enumerate(row):
            if barcode_text(v) and len(barcode_text(v)) >= 12:
                counts[j] += 1
    if not counts:
        return hinted
    top = max(counts.values())
    cands = [j for j, c in counts.items() if c == top]
    return hinted if hinted in cands else cands[0]


def parse_vendor(rows: list[list], sheet: str = "", period: tuple[int, int] | None = None) -> ParseResult:
    """거래처 매입자료 시트를 읽는다.

    날짜는 ① 행마다 있는 일자 열 ② '매입일자 : YYYY-MM-DD' 묶음 머리행
    ③ '[MM월 DD일 … 합계]' 꼬리행(연도는 period) 순으로 찾고, 모두 없으면 날짜 없이 둔다
    (이 경우 대사는 기간 합계 방식으로 한다).
    """
    found = _find_header_row(rows)
    if not found:
        raise ValueError("수량·금액 열 제목을 찾지 못했습니다.")
    hdr_idx, cols = found
    header = rows[hdr_idx]
    start = hdr_idx + 1
    bc_col = _barcode_column(rows, start, cols.get("barcode"))
    if bc_col is None:
        raise ValueError("바코드 열을 찾지 못했습니다.")
    cols["barcode"] = bc_col
    notes: list[str] = []

    if "supply" in cols:
        amount_rule = "공급가액"
    elif "vat" in cols:
        amount_rule = "매입금액 − 부가세액"
    else:
        amount_rule = "매입금액(부가세 포함 추정)"
        notes.append("⚠ 공급가액·부가세 열이 없어 매입금액을 그대로 비교합니다.")

    has_dongwon_supplier = "supplier" in cols and any(
        "동원" in as_text(r[cols["supplier"]]) for r in rows[start:] if cols["supplier"] < len(r)
    )

    recs: list[Record] = []
    pending: list[Record] = []  # 꼬리 합계행이 날짜를 정해 줄 행들
    current: dt.date | None = None
    excluded_supplier = zero_rows = 0
    year = period[0] if period else dt.date.today().year
    for i, row in enumerate(rows[start:], start=start + 1):
        g = lambda f: row[cols[f]] if f in cols and cols[f] < len(row) else None  # noqa: E731
        code = barcode_text(g("barcode"))
        if not code:
            texts = " ".join(as_text(v) for v in row if isinstance(v, str))
            m = _GROUP_DATE_RE.search(texts)
            if m:
                current = dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
                continue
            m = _FOOTER_RE.search(texts)
            if m:
                d = dt.date(year, int(m.group(1)), int(m.group(2)))
                for r in pending:
                    r.date = d
                pending.clear()
            continue

        if has_dongwon_supplier and "동원" not in as_text(g("supplier")):
            excluded_supplier += 1
            continue
        qty = to_number(g("qty")) or 0.0
        if "supply" in cols:
            amt = to_number(g("supply")) or 0.0
        else:
            amt = (to_number(g("gross")) or 0.0) - (to_number(g("vat")) or 0.0)
        if qty == 0 and amt == 0:
            zero_rows += 1
            continue

        date = to_date(g("date")) or current
        name = as_text(g("name"))
        spec = as_text(g("spec"))
        if spec and spec not in name:
            name = f"{name} {spec}"
        rec = Record(date, barcode_key(code), code, name, qty, amt, i, extra={"supplier": as_text(g("supplier"))})
        recs.append(rec)
        if date is None:
            pending.append(rec)

    notes.insert(0, f"거래처 {len(recs):,}행 인식 · 금액 기준: {amount_rule}")
    if excluded_supplier:
        notes.append(f"동원 외 거래처 {excluded_supplier}행 제외")
    if zero_rows:
        notes.append(f"수량·금액 0 행 {zero_rows}건 제외")
    undated = sum(1 for r in recs if r.date is None)
    if recs and undated == len(recs):
        notes.append("날짜 정보가 없는 양식 → 바코드별 기간 합계로 비교")
    elif undated:
        notes.append(f"⚠ 날짜를 정하지 못한 행 {undated}건")
    return ParseResult("vendor", sheet, recs, cols, hdr_idx + 1, notes, list(header), rows[start:], rows)


@dataclass
class Affiliation:
    """거래처 파일에서 감지한 소속(매장) 정보. 못 찾으면 빈 문자열."""

    store: str = ""
    store_source: str = ""  # 어디서 찾았는지(시트 이름 등)
    supplier: str = ""  # 공급처 열의 대표값(예: 동원F&B직거래)

    def short_name(self) -> str:
        """탭·시트 이름에 쓸 짧은 소속명: '(주)세종마트 경산점' → '세종마트 경산점'."""
        return re.sub(r"\(주\)|㈜|주식회사|\(유\)", "", self.store).strip()

    def label(self) -> str:
        parts = []
        if self.store:
            parts.append(f"{self.store} ({self.store_source})")
        if self.supplier:
            parts.append(f"공급처 {self.supplier}")
        return " · ".join(parts)


_BALANCE_SHEET_RE = re.compile(r"^\s*\d{4}\s*(.+?)\s*잔액\s*$")
_STORE_CELL_RE = re.compile(r"(매장|업체|거래처|상호|점포)명?\s*[:：]\s*([^\s:：][^:：]{1,30})$")
_GENERIC_FILE_RE = re.compile(r"^(양식|form|매입|거래처|원장|동원|새|sheet|book|통합 문서)\s*\d*", re.I)


def _norm_store(name: str) -> str:
    n = re.sub(r"\(주\)|㈜|주식회사|\s", "", name)
    return n.replace("마트", "").replace("점", "")


def same_store(a: str, b: str) -> bool:
    """'세종경산'과 '(주)세종마트 경산점'처럼 표기가 달라도 같은 매장으로 본다."""
    x, y = _norm_store(a), _norm_store(b)
    if not x or not y:
        return False
    if x in y or y in x:
        return True
    # 글자 2-gram 겹침 비율
    g = lambda t: {t[i:i + 2] for i in range(len(t) - 1)}  # noqa: E731
    gx, gy = g(x), g(y)
    return bool(gx and gy) and len(gx & gy) / min(len(gx), len(gy)) >= 0.6


def detect_affiliation(path: str, sheets: dict[str, list[list]]) -> Affiliation:
    aff = Affiliation()
    # ① 함께 들어 있는 원장 시트의 판매처명(가장 확실)
    for name, rows in sheets.items():
        if rows and looks_like_ledger(rows):
            try:
                c = Counter(r.customer for r in parse_ledger(rows).records if r.customer)
            except Exception:
                continue
            if c:
                aff.store, aff.store_source = c.most_common(1)[0][0], f"'{name}' 시트 원장"
                break
    # ② 'YYMM 매장명 잔액' 같은 시트 이름
    if not aff.store:
        for name in sheets:
            m = _BALANCE_SHEET_RE.match(name)
            if m:
                aff.store, aff.store_source = m.group(1), f"'{name}' 시트 이름"
                break
    # ③ 위쪽 행의 '매장명 : ○○' 같은 칸
    if not aff.store:
        for name, rows in sheets.items():
            for row in rows[:6]:
                for v in row:
                    m = _STORE_CELL_RE.search(as_text(v)) if isinstance(v, str) else None
                    if m:
                        aff.store, aff.store_source = m.group(2).strip(), f"'{name}' 시트 머리글"
                        break
                if aff.store:
                    break
            if aff.store:
                break
    # ④ 파일 이름(양식1 처럼 일반적인 이름은 제외)
    if not aff.store:
        stem = os.path.splitext(os.path.basename(path))[0]
        if not _GENERIC_FILE_RE.match(stem) and not re.fullmatch(r"[\d_\- ()]+", stem):
            aff.store, aff.store_source = stem, "파일 이름"
    return aff


def inspect_workbook(path: str) -> list[dict]:
    """시트별로 원장/거래처 양식 여부를 판정해 후보 목록을 돌려준다.

    거래처 후보는 날짜가 있고, 공급가(또는 부가세)를 알 수 있고, 행이 많은 순으로 정렬한다.
    각 후보에는 파일에서 감지한 소속(aff)이 함께 들어간다.
    """
    sheets = load_workbook_rows(path)
    aff = detect_affiliation(path, sheets)
    ledgers, vendors = [], []
    for name, rows in sheets.items():
        if not rows:
            continue
        if looks_like_ledger(rows):
            ledgers.append({"sheet": name, "kind": "ledger", "rows": rows, "aff": aff})
            continue
        try:
            res = parse_vendor(rows, name)
        except ValueError:
            continue
        if res.records:
            dated = sum(1 for r in res.records if r.date) / len(res.records)
            exact = bool({"supply", "vat"} & res.columns.keys())
            vendors.append({"sheet": name, "kind": "vendor", "rows": rows, "count": len(res.records),
                            "rank": (dated > 0.5, exact, len(res.records)), "aff": aff})
    vendors.sort(key=lambda c: c["rank"], reverse=True)
    return vendors + ledgers


def supplier_of(vendor: "ParseResult") -> str:
    c = Counter(r.extra.get("supplier") for r in vendor.records if r.extra.get("supplier"))
    return c.most_common(1)[0][0] if c else ""


# ---------------------------------------------------------------------------
# 대사
# ---------------------------------------------------------------------------

STATUSES = ["일치", "수량차이", "단가차이", "동원만", "거래처만"]


@dataclass
class Agg:
    date: dt.date | None
    key: str
    name: str
    qty: float
    amount: float
    rows: list[int]


@dataclass
class ResultRow:
    date: dt.date | None
    key: str
    name: str
    l_qty: float | None
    l_amt: float | None
    v_qty: float | None
    v_amt: float | None
    v_date: dt.date | None = None
    memo: str = ""
    l_rows: list[int] = field(default_factory=list)
    v_rows: list[int] = field(default_factory=list)
    status: str = ""

    @property
    def barcode(self) -> str:
        return display_barcode(self.key)

    @property
    def l_unit(self):
        return self.l_amt / self.l_qty if self.l_qty else None

    @property
    def v_unit(self):
        return self.v_amt / self.v_qty if self.v_qty else None

    @property
    def qty_diff(self) -> float:
        return (self.l_qty or 0) - (self.v_qty or 0)

    @property
    def amt_diff(self) -> float:
        return (self.l_amt or 0) - (self.v_amt or 0)


def _aggregate(recs: list[Record]) -> dict[tuple, Agg]:
    """같은 일자+바코드를 합산한다. 합산 결과 수량·금액이 모두 0이면(입고 후 당일 반품 등) 뺀다."""
    out: dict[tuple, Agg] = {}
    for r in recs:
        k = (r.date, r.key)
        a = out.get(k)
        if a is None:
            out[k] = Agg(r.date, r.key, r.name, r.qty, r.amount, [r.row])
        else:
            a.qty += r.qty
            a.amount += r.amount
            a.rows.append(r.row)
    return {k: a for k, a in out.items() if abs(a.qty) > 1e-9 or abs(a.amount) > 1e-6}


def _memo_single(qty: float, amount: float, side: str) -> str:
    if abs(qty) < 1e-9:
        return "수량 없이 금액만(정산/조정성)"
    if abs(amount) < 1e-6:
        return "무상/판촉물(금액 0)"
    if qty < 0:
        return f"{side} 반품/마이너스"
    return ""


def classify(row: ResultRow, tol: float) -> str:
    if row.l_qty is None:
        return "거래처만"
    if row.v_qty is None:
        return "동원만"
    if abs(row.qty_diff) > 1e-9:
        return "수량차이"
    return "일치" if abs(row.amt_diff) <= tol else "단가차이"


def reconcile(ledger: list[Record], vendor: list[Record], window_days: int = 3, tolerance: float = 1.0) -> list[ResultRow]:
    """일자+바코드로 맞추고, 남은 건은 같은 바코드끼리 ±window_days 안에서 짝짓는다."""
    la, va = _aggregate(ledger), _aggregate(vendor)
    rows: list[ResultRow] = []
    l_only: dict[str, list[Agg]] = defaultdict(list)
    v_only: dict[str, list[Agg]] = defaultdict(list)
    for k, a in la.items():
        b = va.get(k)
        if b is None:
            l_only[a.key].append(a)
        else:
            rows.append(ResultRow(a.date, a.key, a.name, a.qty, a.amount, b.qty, b.amount, l_rows=a.rows, v_rows=b.rows))
    for k, b in va.items():
        if k not in la:
            v_only[b.key].append(b)

    for key, ls in l_only.items():
        # 원장 날짜순으로, 아직 짝이 없는 거래처 건 중 가장 가까운 날짜를 가져간다
        # (동률이면 거래처 날짜가 이른 쪽 — 기존 수작업 대사 결과와 같은 규칙).
        vs = sorted(v_only.get(key, []), key=lambda b: b.date or dt.date.max)
        used_v: set[int] = set()
        for a in sorted(ls, key=lambda a: a.date or dt.date.max):
            best = None
            for j, b in enumerate(vs):
                if j in used_v or not (a.date and b.date):
                    continue
                diff = (b.date - a.date).days
                if abs(diff) <= window_days and (best is None or (abs(diff), diff) < (abs(best[1]), best[1])):
                    best = (j, diff)
            if best is None:
                rows.append(ResultRow(a.date, key, a.name, a.qty, a.amount, None, None, l_rows=a.rows,
                                      memo=_memo_single(a.qty, a.amount, "동원")))
                continue
            j, diff = best
            used_v.add(j)
            b = vs[j]
            rows.append(
                ResultRow(a.date, key, a.name, a.qty, a.amount, b.qty, b.amount, v_date=b.date,
                          memo=f"거래처 일자 {diff:+d}일", l_rows=a.rows, v_rows=b.rows)
            )
        v_only[key] = [b for j, b in enumerate(vs) if j not in used_v]

    for key, vs in v_only.items():
        for b in vs:
            rows.append(ResultRow(b.date, key, b.name, None, None, b.qty, b.amount, v_rows=b.rows,
                                  memo=_memo_single(b.qty, b.amount, "거래처")))

    for r in rows:
        r.status = classify(r, tolerance)
    rows.sort(key=lambda r: (r.date or dt.date.max, r.key))
    return rows


@dataclass
class Summary:
    l_qty: float
    v_qty: float
    l_amt: float
    v_amt: float
    counts: dict[str, int]

    @property
    def total(self) -> int:
        return sum(self.counts.values())

    @property
    def to_check(self) -> int:
        """사람이 확인해야 할 줄(일치가 아닌 줄) 수."""
        return self.total - self.counts.get("일치", 0)

    @property
    def match_rate(self) -> float:
        return self.counts.get("일치", 0) / self.total if self.total else 0.0


def summarize(rows: list[ResultRow]) -> Summary:
    c = Counter(r.status for r in rows)
    return Summary(
        l_qty=sum(r.l_qty or 0 for r in rows),
        v_qty=sum(r.v_qty or 0 for r in rows),
        l_amt=sum(r.l_amt or 0 for r in rows),
        v_amt=sum(r.v_amt or 0 for r in rows),
        counts={s: c.get(s, 0) for s in STATUSES},
    )


def ledger_period(recs: list[Record]) -> tuple[int, int] | None:
    c = Counter((r.date.year, r.date.month) for r in recs if r.date)
    return c.most_common(1)[0][0] if c else None


def best_customer(ledger: list[Record], vendor: list[Record]) -> str | None:
    """거래처 자료와 (일자, 바코드)가 가장 많이 겹치는 원장 판매처를 고른다."""
    vkeys = {(r.date, r.key) for r in vendor}
    vbars = {r.key for r in vendor}
    score: Counter = Counter()
    for r in ledger:
        score[r.customer] += 3 if (r.date, r.key) in vkeys else (1 if r.key in vbars else 0)
    if not score:
        return None
    return score.most_common(1)[0][0]


# ---------------------------------------------------------------------------
# 한 번의 비교 작업(원장 1 × 거래처 파일 1)
# ---------------------------------------------------------------------------


@dataclass
class CompareResult:
    label: str  # 결과 시트 이름에 쓰는 짧은 이름
    vendor_file: str
    vendor_sheet: str
    customer: str | None
    mode: str  # "daily" | "period"
    rows: list[ResultRow]
    summary: Summary
    notes: list[str]
    vendor: ParseResult
    ledger_rows_used: list[Record]
    warning: str = ""  # 자료 짝이 맞지 않아 보일 때 화면 상단에 띄울 경고
    affiliation: "Affiliation" = field(default_factory=lambda: Affiliation())


def compare(
    ledger: ParseResult,
    vendor: ParseResult,
    *,
    label: str,
    vendor_file: str = "",
    customer: str | None = "auto",
    window_days: int = 3,
    tolerance: float = 1.0,
    clip_to_ledger: bool = True,
    affiliation: Affiliation | None = None,
) -> CompareResult:
    """원장과 거래처 자료 하나를 대사한다.

    customer: "auto" 면 겹치는 건이 가장 많은 판매처, None 이면 원장 전체.
    clip_to_ledger: 원장 기간(±window) 밖의 거래처 행은 빼고 비교한다.
    """
    notes = list(vendor.notes)
    lrecs = ledger.records
    customers = {r.customer for r in lrecs}
    if customer == "auto":
        customer = best_customer(lrecs, vendor.records) if len(customers) > 1 else None
    if customer:
        lrecs = [r for r in lrecs if r.customer == customer]
        if len(customers) > 1:
            notes.append(f"원장 판매처: {customer} ({len(lrecs):,}행)")

    vrecs = vendor.records
    dated = [r for r in vrecs if r.date]
    mode = "daily" if len(dated) >= len(vrecs) * 0.5 else "period"
    if mode == "daily":
        if len(dated) < len(vrecs):
            notes.append(f"⚠ 날짜 없는 거래처 {len(vrecs) - len(dated)}행은 비교에서 뺐습니다")
        vrecs = dated
        ldates = [r.date for r in lrecs if r.date]
        if clip_to_ledger and ldates:
            lo = min(ldates) - dt.timedelta(days=window_days)
            hi = max(ldates) + dt.timedelta(days=window_days)
            inside = [r for r in vrecs if lo <= r.date <= hi]
            if len(inside) < len(vrecs):
                notes.append(
                    f"원장 기간({min(ldates):%m/%d}~{max(ldates):%m/%d}) 밖 거래처 {len(vrecs) - len(inside):,}행 제외"
                )
            vrecs = inside
        rows = reconcile(lrecs, vrecs, window_days, tolerance)
    else:
        # 날짜가 없는 양식: 원장도 날짜를 지우고 바코드별 합계로 비교한다.
        strip = lambda rs: [Record(None, r.key, r.barcode, r.name, r.qty, r.amount, r.row, r.customer) for r in rs]  # noqa: E731
        rows = reconcile(strip(lrecs), strip(vrecs), 0, tolerance)

    if not vrecs:
        notes.append("⚠ 비교할 거래처 행이 없습니다(기간·판매처를 확인하세요)")
    summary = summarize(rows)
    aff = Affiliation(**vars(affiliation)) if affiliation else Affiliation()
    aff.supplier = aff.supplier or supplier_of(vendor)
    who = customer or (lrecs[0].customer if lrecs else "")
    warning = mismatch_warning(lrecs, vrecs, summary, customer)
    # 파일 이름으로만 짐작한 소속은 경고 근거로 쓰지 않는다(표시만).
    if aff.store and aff.store_source != "파일 이름" and who and not same_store(aff.store, who):
        warning = (f"거래처 자료는 '{aff.store}'({aff.store_source})로 보이는데 원장 판매처는 '{who}'입니다. "
                   + (warning or "같은 매장 자료인지 확인해 주세요."))
    return CompareResult(label, vendor_file, vendor.sheet, customer, mode, rows, summary, notes, vendor, lrecs,
                         warning, aff)


def mismatch_warning(lrecs: list[Record], vrecs: list[Record], summary: Summary, customer: str | None) -> str:
    """맞는 줄이 거의 없으면(5% 미만), 날짜를 빼고 바코드별 수량 합계로도 맞는지 봐서 '다른 매장 자료' 가능성을 알린다."""
    if summary.total < 20 or summary.match_rate >= 0.05:
        return ""
    lq: dict[str, float] = defaultdict(float)
    vq: dict[str, float] = defaultdict(float)
    for r in lrecs:
        lq[r.key] += r.qty
    for r in vrecs:
        vq[r.key] += r.qty
    common = set(lq) & set(vq)
    same = sum(1 for k in common if abs(lq[k] - vq[k]) < 1e-9)
    who = customer or (lrecs[0].customer if lrecs else "")
    return (
        f"{summary.total:,}줄 중 맞는 줄이 {summary.counts.get('일치', 0)}줄뿐입니다. 날짜를 무시하고 바코드별 수량 합계로 봐도 "
        f"{len(vq)}종 중 {same}종만 같아, 원장 판매처({who})와 다른 매장의 자료일 가능성이 큽니다. "
        "같은 매장의 원장인지, 기간이 같은지 확인해 주세요."
    )
